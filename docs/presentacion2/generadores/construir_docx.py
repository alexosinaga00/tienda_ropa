"""Genera FashionStore_Presentacion2_Ciclo2.docx a partir del documento de la
Presentacion 1 (Ciclo 1), aplicando:

  1. las correcciones de texto donde el documento contradice al codigo,
  2. las tablas actualizadas (30 CU, priorizacion por ciclos, paquetes),
  3. las 16 fichas nuevas del Ciclo 2 con su diagrama particular,
  4. los diagramas nuevos y corregidos exportados de Enterprise Architect
     (comunicacion, clases de analisis, secuencia, UC-02, VP, PKG, DESP).

El original nunca se modifica: se parte de una copia.
"""
from __future__ import annotations

import copy
import json
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Emu, Inches
from docx.oxml.ns import qn
from PIL import Image

AQUI = Path(__file__).parent
DIAG = AQUI / "diagramas"
ORIGEN = AQUI.parent / "FashionStore_Presentacion2_Ciclo1.docx"
DESTINO = AQUI.parent / "FashionStore_Presentacion2_Ciclo2.docx"

shutil.copy(ORIGEN, DESTINO)
doc = Document(str(DESTINO))
body = doc.element.body
bloques = list(body.iterchildren())


# --------------------------------------------------------------- utilidades
def texto(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _es_negrita(run) -> bool:
    rpr = run.find(qn("w:rPr"))
    return rpr is not None and rpr.find(qn("w:b")) is not None


def poner_texto(p, nuevo: str, negrita: bool | None = None) -> None:
    """Reescribe el parrafo conservando su pPr y el formato del primer run.

    Si el original abre con un tramo en negrita seguido de texto normal (viñetas
    del tipo «Título: explicación»), se respeta ese patrón hasta el primer ':'.
    negrita=False fuerza texto normal."""
    runs = [r for r in p.findall(qn("w:r")) if r.find(qn("w:t")) is not None] or p.findall(qn("w:r"))
    if not runs:
        r = p.makeelement(qn("w:r"), {})
        p.append(r)
        runs = [r]
    primero = runs[0]
    normal = next((r for r in runs[1:] if not _es_negrita(r)), None)
    patron_titulo = _es_negrita(primero) and normal is not None and ":" in nuevo[:90]
    modelo_normal = copy.deepcopy(normal) if normal is not None else None
    for r in p.findall(qn("w:r")):
        if r is not primero:
            p.remove(r)
    for hijo in list(primero):
        if hijo.tag != qn("w:rPr"):
            primero.remove(hijo)
    if negrita is False:
        rpr = primero.find(qn("w:rPr"))
        if rpr is not None and rpr.find(qn("w:b")) is not None:
            rpr.remove(rpr.find(qn("w:b")))
    cabeza, cola = nuevo, ""
    if patron_titulo and negrita is None:
        corte = nuevo.index(":") + 1
        cabeza, cola = nuevo[:corte], nuevo[corte:]
    t = primero.makeelement(qn("w:t"), {})
    t.text = cabeza
    t.set(qn("xml:space"), "preserve")
    primero.append(t)
    if cola:
        segundo = modelo_normal
        for hijo in list(segundo):
            if hijo.tag != qn("w:rPr"):
                segundo.remove(hijo)
        t2 = segundo.makeelement(qn("w:t"), {})
        t2.text = cola
        t2.set(qn("xml:space"), "preserve")
        segundo.append(t2)
        primero.addnext(segundo)


def parrafos_de_celda(celda):
    return celda.findall(qn("w:p"))


def poner_celda(celda, lineas: list[str]) -> None:
    """Deja en la celda una linea por elemento, clonando el formato del primer parrafo."""
    ps = parrafos_de_celda(celda)
    plantilla = copy.deepcopy(ps[0])
    for p in ps[1:]:
        celda.remove(p)
    poner_texto(ps[0], lineas[0])
    anterior = ps[0]
    for linea in lineas[1:]:
        nuevo = copy.deepcopy(plantilla)
        poner_texto(nuevo, linea)
        anterior.addnext(nuevo)
        anterior = nuevo


def filas(tabla):
    return tabla.findall(qn("w:tr"))


def celdas(fila):
    return fila.findall(qn("w:tc"))


def fila_por_etiqueta(tabla, etiqueta: str):
    for f in filas(tabla):
        cs = celdas(f)
        if cs and texto(cs[0]).strip().upper().startswith(etiqueta.upper()):
            return f
    raise KeyError(etiqueta)


def poner_ficha(tabla, campos: dict[str, list[str]]) -> None:
    for etiqueta, lineas in campos.items():
        f = fila_por_etiqueta(tabla, etiqueta)
        poner_celda(celdas(f)[1], lineas)


# --- imagenes ---------------------------------------------------------------
_scratch = doc.add_paragraph()  # parrafo auxiliar al final, se borra al terminar


def parrafo_imagen(archivo: Path, ancho_max_in: float = 6.29, alto_max_in: float = 8.2, centrado=True):
    """Crea un parrafo nuevo (todavia sin ubicar) con la imagen escalada."""
    with Image.open(archivo) as im:
        w, h = im.size
    ancho = ancho_max_in
    if h / w * ancho > alto_max_in:
        ancho = alto_max_in * w / h
    p = doc.add_paragraph()
    p.add_run().add_picture(str(archivo), width=Inches(ancho))
    if centrado:
        ppr = p._p.get_or_add_pPr()
        jc = ppr.makeelement(qn("w:jc"), {qn("w:val"): "center"})
        ppr.append(jc)
    el = p._p
    body.remove(el)
    return el


def reemplazar_imagen(p, archivo: Path, ancho_max_in: float = 6.29, alto_max_in: float = 8.2) -> None:
    """Sustituye la imagen de un parrafo existente por otra, recalculando el tamano."""
    nuevo = parrafo_imagen(archivo, ancho_max_in, alto_max_in)
    p.addnext(nuevo)
    p.getparent().remove(p)


def insertar_despues(ref, elementos: list):
    anterior = ref
    for el in elementos:
        anterior.addnext(el)
        anterior = el
    return anterior


def parrafo_como(modelo, texto_nuevo: str, negrita: bool | None = None):
    nuevo = copy.deepcopy(modelo)
    poner_texto(nuevo, texto_nuevo, negrita)
    return nuevo


# --- numeracion propia por ficha (para que el flujo arranque en 1) ----------
numbering = doc.part.numbering_part.element
_abstract = None
for num in numbering.findall(qn("w:num")):
    if num.get(qn("w:numId")) == "23":
        _abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
_siguiente_num = [500]


def nuevo_num_id() -> str:
    _siguiente_num[0] += 1
    nid = str(_siguiente_num[0])
    num = numbering.makeelement(qn("w:num"), {qn("w:numId"): nid})
    ab = num.makeelement(qn("w:abstractNumId"), {qn("w:val"): _abstract})
    num.append(ab)
    lvl = num.makeelement(qn("w:lvlOverride"), {qn("w:ilvl"): "0"})
    start = lvl.makeelement(qn("w:startOverride"), {qn("w:val"): "1"})
    lvl.append(start)
    num.append(lvl)
    numbering.append(num)
    return nid


def renumerar_flujo(tabla) -> None:
    nid = nuevo_num_id()
    f = fila_por_etiqueta(tabla, "FLUJO PRINCIPAL")
    for p in parrafos_de_celda(celdas(f)[1]):
        numpr = p.find(qn("w:pPr"))
        if numpr is None:
            continue
        np = numpr.find(qn("w:numPr"))
        if np is not None:
            np.find(qn("w:numId")).set(qn("w:val"), nid)


# =========================================================== 1. TEXTOS
CAMBIOS_TEXTO = [
    ("Este documento corresponde a la primera de tres presentaciones",
     "Este documento corresponde a la segunda de tres presentaciones del examen de Sistemas de "
     "Información II y se construye sobre el entregado en la Presentación 1. Para mantener la "
     "profundidad de análisis y diseño que exige la metodología en un plazo acotado, el trabajo se "
     "organiza en ciclos: el Ciclo 1 (14 casos de uso) cubrió el flujo núcleo de punta a punta "
     "—identidad y acceso, catálogo, disponibilidad de inventario, reserva multi-prenda, vestidor "
     "virtual y venta con pago digital y presencial— y esta entrega agrega el Ciclo 2 (16 casos de "
     "uso): administración de usuarios y roles, ciudades y empleados, búsqueda con filtros, "
     "abastecimiento (proveedores y recepción de mercadería), inventario global y sus movimientos, "
     "el ciclo completo de la reserva desde el cliente y desde la sucursal, carrito, historial de "
     "compras, promociones, ventas por sucursal y cobro en caja. Con eso el documento llega a 30 de "
     "los 34 casos de uso del catálogo vigente del sistema; los 4 restantes (confirmación de "
     "transacción de pasarela, recomendador, reportes y dashboard) y las funcionalidades "
     "complementarias ya implementadas se documentan en la Presentación final."),
    ("Desarrollar e implementar el núcleo transaccional de FashionStore",
     "Desarrollar e implementar la plataforma FashionStore — identidad y control de acceso, "
     "organización multi-sucursal, catálogo por variante talla-color, abastecimiento, inventario por "
     "sucursal, reserva de prendas previa a la compra, vestidor virtual por realidad aumentada, y "
     "venta digital y presencial con pasarelas de pago en sandbox — aplicando el Proceso Unificado "
     "de Desarrollo de Software y modelado UML 2.5 de principio a fin, documentada por ciclos: 14 "
     "casos de uso en el Ciclo 1 y 16 en el Ciclo 2, sobre un catálogo total de 34."),
    ("El alcance de este ciclo contempla los siguientes módulos:",
     "El alcance documentado hasta esta entrega (Ciclos 1 y 2) contempla los siguientes módulos:"),
    ("Identidad y control de acceso: registro de clientes e inicio de sesión con roles",
     "Identidad y control de acceso: registro de clientes, inicio de sesión con roles y "
     "administración de usuarios, roles y permisos"),
    ("Organización: gestión de sucursales",
     "Organización: gestión de ciudades, sucursales con sus horarios de atención, y empleados "
     "asignados a una sucursal"),
    ("Catálogo: gestión de catálogo maestro incluyendo tabla de medidas",
     "Catálogo: gestión del catálogo maestro (categorías, tallas, colores, materiales, temporadas y "
     "colecciones), gestión de productos con variantes, tabla de medidas e imágenes, consulta del "
     "catálogo, búsqueda con filtros, y carga de assets y anclajes para el vestidor virtual"),
    ("Inventario: consulta de disponibilidad por sucursal",
     "Abastecimiento: gestión de proveedores y registro de la recepción de mercadería, que es la "
     "entrada de stock con costo"),
    ("Reservas: reserva de múltiples prendas y atención de la prueba en sucursal",
     "Inventario: consulta de disponibilidad por sucursal, inventario consolidado global, y registro "
     "de movimientos (ajustes y transferencias entre sucursales)"),
    ("Vestidor virtual: prueba de prenda en modo espejo por realidad aumentada",
     "Reservas: reserva de múltiples prendas, consulta y cancelación por parte del cliente, y "
     "consulta y atención de la prueba en la sucursal"),
    ("Ventas y pagos: compra digital con pasarela, venta presencial en punto de caja, y pago med",
     "Vestidor virtual: prueba de prenda en modo espejo por realidad aumentada, con assets y "
     "anclajes cargados desde el back office"),
    ("Dirigido por casos de uso: los 14 CU del Ciclo 1 son el hilo conductor",
     "Dirigido por casos de uso: los 30 CU documentados (14 del Ciclo 1 y 16 del Ciclo 2) son el "
     "hilo conductor de requisitos, análisis, diseño e implementación."),
    ("Iterativo e incremental: el proyecto completo se organiza en tres presentaciones",
     "Iterativo e incremental: el proyecto se organiza en tres presentaciones que amplían "
     "progresivamente el alcance documentado; cada ciclo atraviesa completo los flujos de "
     "requisitos, análisis y diseño antes de agregar el siguiente."),
    ("Inicio: definición del alcance del Ciclo 1 y viabilidad técnica",
     "Inicio: definición del alcance de los ciclos y viabilidad técnica del stack (FastAPI + "
     "PostgreSQL + Flutter/Angular)."),
    ("Construcción: desarrollo iterativo de los 14 casos de uso del Ciclo 1",
     "Construcción: desarrollo iterativo de los casos de uso; a la fecha de este documento los 30 "
     "casos de uso de los Ciclos 1 y 2 están implementados en el backend y con pantalla en la web "
     "(back office y tienda pública) o en la aplicación móvil, según su actor (ver 4.2)."),
    ("Esta presentación cubre Captura de Requisitos, Análisis y Diseño (hasta Diseño de Datos)",
     "Esta presentación cubre Captura de Requisitos, Análisis y Diseño (hasta Diseño de Datos) para "
     "los Ciclos 1 y 2, e Implementación (F.T.4) para el sistema construido. El flujo de Pruebas "
     "sobre casos de uso y los diagramas de estado, tiempo y navegación se documentan en la "
     "Presentación final, junto con los 4 casos de uso restantes del catálogo y las "
     "funcionalidades complementarias."),
    ("Exige coordinar en el mismo ciclo un caso de uso técnicamente exigente",
     "Exige coordinar en el mismo ciclo un caso de uso técnicamente exigente (el vestidor virtual) "
     "junto con el resto del núcleo transaccional."),
    ("El recorte a 14 CU obliga a documentar explícitamente qué queda fuera",
     "Documentar por ciclos obliga a declarar explícitamente qué queda fuera de cada entrega y en "
     "qué presentación se retoma, y a revisar en cada ciclo lo ya escrito para que siga coincidiendo "
     "con el código."),
    ("Borrado lógico y auditoría transversal (creado_por, creado_en, activo) en toda tabla de ne",
     "Borrado lógico donde la entidad se da de baja sin perder historial (campo activo) y columnas "
     "de auditoría (creado_por, creado_en) en las tablas cuyo ciclo de vida las necesita."),
    ("Se verifica que la partición en 9 paquetes mantiene bajo acoplamiento",
     "Se verifica que la partición en paquetes mantiene bajo acoplamiento (dependencias "
     "unidireccionales, sin ciclos entre paquetes de negocio, ningún paquete accede a las tablas de "
     "otro) y alta cohesión. Los Ciclos 1 y 2 involucran 10 paquetes: los 9 del Ciclo 1 más "
     "abastecimiento; el backend tiene además entregas, inteligencia y reportes, que corresponden a "
     "casos de uso de la Presentación final o quedan fuera del alcance documentado."),
    ("Las 33 clases de datos (agrupadas por dominio",
     "El diagrama DCD muestra las entidades clave del núcleo transaccional (usuario y cliente, "
     "empleado y sucursal, categoría, producto y variante, stock, reserva, carrito, venta y su "
     "detalle, y pago); el esquema completo está en los diagramas DAT-01 a DAT-06 y en "
     "docs/fashionstore_esquema.sql. Las 60 tablas del esquema están en Tercera Forma Normal: cada "
     "atributo no clave depende únicamente de la clave primaria de su tabla, y no existen grupos "
     "repetitivos (por eso reserva_detalle, venta_detalle y carrito_detalle son tablas propias, no "
     "columnas repetidas dentro de reserva, venta o carrito). Las relaciones N:M sin atributos "
     "propios (usuario-rol, rol-permiso) se modelan como asociación directa, sin clase intermedia, "
     "siguiendo la notación UML 2.5 estándar. producto_variante.id, nunca producto.id, es la clave "
     "foránea real en stock, reserva_detalle, venta_detalle y carrito_detalle: es la decisión de "
     "diseño más importante del modelo, porque el producto es el concepto comercial y la variante es "
     "la unidad real de negocio."),
    ("El esquema físico completo y definitivo (61 tablas, 2 vistas)",
     "El esquema físico de referencia (61 tablas y 2 vistas) está documentado en "
     "docs/fashionstore_esquema.sql; es documentación, no se ejecuta contra la base — las tablas las "
     "crea Alembic a partir de los modelos SQLAlchemy. La base implementada tiene 60 de esas tablas "
     "y las 2 vistas: la única que no se creó es bitacora, porque la auditoría de eventos no "
     "corresponde a ningún caso de uso de los ciclos documentados. De las ocho decisiones de diseño "
     "de datos del proyecto, las que aplican directamente al núcleo son:"),
    ("Auditoría y borrado lógico transversales: toda tabla de negocio incluye creado_por, creado",
     "Auditoría y borrado lógico: el borrado es siempre lógico (activo = false), nunca DELETE "
     "físico, en toda tabla de negocio que conserve historial; las columnas de auditoría se aplican "
     "según lo que necesita el ciclo de vida de cada tabla (producto lleva creado_por, creado_en y "
     "activo; producto_variante y empleado solo activo; los movimientos de inventario y las "
     "reservas llevan creado_en y usuario). La decisión fija el criterio, cada tabla toma lo que le "
     "corresponde."),
    # --- F.T.4 ---
    ("API: https://tiendaropa-production-b36a.up.railway.app",
     "API: https://tiendaropa-production-b36a.up.railway.app — GET /health responde 200. El "
     "despliegue alterno de la misma aplicación en "
     "https://backend-production-0714.up.railway.app también responde 200."),
    ("La API expuesta suma 193 operaciones HTTP sobre 126 rutas, agrupadas en 38 etiquetas",
     "La API expuesta suma 193 operaciones HTTP sobre 126 rutas, agrupadas en 38 etiquetas "
     "funcionales, y la suite de pruebas del backend tiene 234 funciones de prueba. El esquema de "
     "datos consta de 60 tablas base y 2 vistas, creadas por 17 migraciones de Alembic: una tabla "
     "menos que la especificación de referencia docs/fashionstore_esquema.sql, que documenta 61 "
     "(la tabla bitacora no se implementó; ver 6.3.2).", 3),
    ("Una auditoría interna del código realizada el 09/09/2026 sobre las tres capas",
     "Una auditoría interna del código realizada el 09/09/2026 sobre las tres capas "
     "(docs/auditoria_2026-09-09.txt) registró tres hallazgos. Los tres ya fueron corregidos en el "
     "código y se documentan acá junto con lo que sigue abierto:", 2),
    ("Condición de carrera en la resolución de pagos (pagos/service.py): la idempotencia está pe",
     "Condición de carrera en la resolución de pagos: pagos/service.py ahora obtiene la fila del "
     "pago y la de la venta con bloqueo (SELECT FOR UPDATE, vía pago_repo.obtener_bloqueado y "
     "venta_repo.obtener_bloqueada) antes de resolver el resultado de la pasarela, igual que "
     "inventario.registrar_movimiento(). Dos confirmaciones concurrentes del mismo pago ya no "
     "pueden duplicar su efecto."),
    ("Falta de guardia contra dos pagos activos sobre la misma venta: un doble clic o un reinten",
     "Dos pagos activos sobre la misma venta: iniciar_pago_pasarela() exige que la venta esté "
     "«pendiente de pago» y que no exista ya un pago en estado «iniciado», así que un doble clic o "
     "un reintento de red no abre un segundo pago."),
    ("Desvío del patrón por capas en varios routers: algunos routers llaman directamente al repo",
     "Desvío del patrón por capas: ningún router importa ya un repositorio; todos pasan por su "
     "service, y seguridad.service lanza excepciones de dominio en vez de HTTPException."),
    ("Los dos primeros son los candidatos naturales a resolver en la siguiente iteración, por",
     "Sigue abierto, y se documenta como deuda técnica: pagar_en_caja() no verifica que no haya un "
     "pago de pasarela en curso para la misma venta (se apoya solo en el bloqueo y la validación de "
     "estado de confirmar_venta()); el permiso probador.usar existe en la base pero ningún endpoint "
     "lo verifica, le alcanza con la sesión; el detalle público de producto devuelve "
     "cantidad_disponible en nulo y la disponibilidad se consulta con el endpoint de inventario "
     "(CU-13); y la expiración de reservas depende de un cron externo que invoque "
     "/tareas/expirar-reservas, porque el proyecto no incluye un planificador propio.", 1),
    ("__NO_USAR__",
     "Sigue abierto, y se documenta como deuda técnica: pagar_en_caja() no verifica que no haya un "
     "pago de pasarela en curso para la misma venta (se apoya solo en el bloqueo y la validación de "
     "estado de confirmar_venta()); el permiso probador.usar existe en la base pero ningún endpoint "
     "lo verifica, le alcanza con la sesión; el detalle público de producto devuelve "
     "cantidad_disponible en nulo y la disponibilidad se consulta con el endpoint de inventario "
     "(CU-13); y la expiración de reservas depende de un cron externo que invoque "
     "/tareas/expirar-reservas, porque el proyecto no incluye un planificador propio."),
    ("Modo espejo (obligatorio, exclusivo de la app Flutter). Usa google_mlkit_pose_detection so",
     "Modo espejo (obligatorio, exclusivo de la app Flutter). Usa google_mlkit_pose_detection sobre "
     "la cámara del dispositivo: detecta la pose del usuario y superpone el overlay de la prenda "
     "alineándolo con los puntos de anclaje marcados previamente por el administrador (CU-21). El "
     "procesamiento ocurre en el teléfono; el backend solo provee el asset validado y registra la "
     "sesión."),
    ("protocolos involucrados ya están representados en el diagrama de despliegue DIS-01, y",
     "protocolos involucrados ya están representados en el diagrama de despliegue DESP-01, y"),
]

parrafos = [b for b in bloques if b.tag == qn("w:p")]
usados = set()
for entrada in CAMBIOS_TEXTO:
    prefijo, nuevo = entrada[0], entrada[1]
    borrar = entrada[2] if len(entrada) > 2 else 0
    if prefijo == "__NO_USAR__":
        continue
    for p in parrafos:
        if id(p) in usados:
            continue
        if texto(p).strip().startswith(prefijo[:60]):
            poner_texto(p, nuevo)
            usados.add(id(p))
            # párrafos que eran continuación de línea del mismo texto (F.T.4 viene cortado por renglón)
            for _ in range(borrar):
                sig = p.getnext()
                if sig is not None and sig.tag == qn("w:p"):
                    usados.add(id(sig))
                    body.remove(sig)
            break
    else:
        print("  [OJO] no se encontró el párrafo:", prefijo[:60])

# el alcance tenía 7 módulos y ahora son 8: se agrega ventas y pagos al final de la lista
for p in list(body.iterchildren()):
    if texto(p).strip().startswith("Vestidor virtual: prueba de prenda en modo espejo por realidad aumentada, con assets"):
        extra = parrafo_como(p, "Ventas y pagos: carrito con las promociones vigentes, compra digital con "
                                "pasarela, venta presencial con cobro en caja, historial de compras del "
                                "cliente, ventas por sucursal y gestión de promociones")
        p.addnext(extra)
        break

# títulos y numeración del perfil / marco teórico
RENOMBRES = {
    "1.2 Objetivos": "1.2 Objetivo General",
    "Objetivo General": None,          # se elimina el subtítulo, ahora redundante
    "Objetivos Específicos": "1.3 Objetivos Específicos",
    "1.3 Descripción del problema": "1.4 Descripción del problema",
    "1.4 Alcance": "1.5 Alcance",
    "2.1.1 Fundamentos": "2.1 Comercio electrónico multi-sucursal: fundamentos",
    "4.1 Selección de plataforma de software": "7.1 Selección de plataforma de software",
    "4.2 Implementación de la arquitectura del sistema principal":
        "7.2 Implementación de la arquitectura del sistema principal",
    "4.3 Implementación de la arquitectura del subsistema":
        "7.3 Implementación de la arquitectura del subsistema",
}
for p in parrafos:
    t = texto(p).strip()
    if t in RENOMBRES:
        destino_txt = RENOMBRES[t]
        if destino_txt is None:
            if t == "Objetivo General":
                p.getparent().remove(p)
        else:
            poner_texto(p, destino_txt)
        RENOMBRES.pop(t)

# el nivel de título de "1.3 Objetivos Específicos" pasa de Ttulo3 a Ttulo2
for p in parrafos:
    if texto(p).strip() == "1.3 Objetivos Específicos":
        ppr = p.find(qn("w:pPr"))
        est = ppr.find(qn("w:pStyle"))
        if est is not None:
            est.set(qn("w:val"), "Ttulo2")

for p in list(body.iterchildren()):
    if texto(p).strip().startswith("Integrar pasarelas de pago en modo sandbox (Libélula y PayPal)"):
        extra = parrafo_como(p, "Implementar la administración del back office que sostiene la operación "
                                "diaria: usuarios, roles y permisos; ciudades y empleados por sucursal; "
                                "proveedores y recepción de mercadería con costo promedio ponderado; "
                                "inventario consolidado y sus movimientos; y promociones.")
        p.addnext(extra)
        break

print("Textos aplicados")

cus = json.loads((AQUI / "cu_ciclo2.json").read_text(encoding="utf8"))
CICLO1 = [
    ("CU-01", "Registrar cliente", "Crear una cuenta de cliente para acceder al catálogo, reservar, comprar y usar el vestidor virtual.", "Cliente", "Web y App móvil", "seguridad"),
    ("CU-02", "Iniciar sesión", "Autenticarse para acceder a las funciones del rol propio.", "Cliente, Administrador, Encargado, Cajero, Proveedor", "Web y App móvil", "seguridad"),
    ("CU-05", "Gestionar sucursales", "Mantener las sucursales físicas y sus horarios de atención.", "Administrador", "Web (back office)", "organizacion"),
    ("CU-07", "Gestionar catálogo maestro", "Mantener categorías, tallas, colores, materiales, temporadas y colecciones.", "Administrador", "Web (back office)", "catalogo"),
    ("CU-08", "Gestionar productos (variantes e imágenes)", "Crear productos, variantes talla-color, tabla de medidas e imágenes.", "Administrador", "Web (back office)", "catalogo"),
    ("CU-09", "Consultar catálogo", "Navegar el catálogo de productos disponibles.", "Cliente", "Web y App móvil", "catalogo"),
    ("CU-13", "Consultar disponibilidad por sucursal", "Conocer en qué sucursal hay disponibilidad física de una variante.", "Cliente", "Web y App móvil", "inventario"),
    ("CU-16", "Reservar varias prendas", "Reservar una o varias variantes en una sucursal, con fecha y franja horaria, antes de comprar.", "Cliente", "Web y App móvil", "reservas"),
    ("CU-20", "Atender prueba de reserva en sucursal", "Gestionar el ciclo completo de atención de una reserva.", "Encargado de sucursal", "Web (back office)", "reservas"),
    ("CU-21", "Cargar assets y marcar anclajes", "Habilitar una variante para el vestidor virtual.", "Administrador", "Web (back office)", "probador"),
    ("CU-22", "Probar prenda en modo espejo", "Ver la prenda superpuesta en tiempo real mediante realidad aumentada.", "Cliente", "App móvil", "probador"),
    ("CU-24", "Realizar compra digital", "Completar una compra con pago mediante pasarela electrónica.", "Cliente", "Web y App móvil", "ventas"),
    ("CU-25", "Registrar venta presencial", "Registrar en punto de caja la venta de las prendas decididas.", "Cajero", "Web (back office)", "ventas"),
    ("CU-29", "Pagar mediante pasarela digital", "Pagar una venta digital a través de una pasarela de pago.", "Cliente", "Web y App móvil", "pagos"),
]
PRIORIDAD_C1 = {"CU-01": ("Crítica", "Bajo", "Cliente"), "CU-02": ("Crítica", "Alto", "Todos"),
                "CU-05": ("Alta", "Bajo", "Administrador"), "CU-07": ("Alta", "Bajo", "Administrador"),
                "CU-08": ("Alta", "Medio", "Administrador"), "CU-09": ("Crítica", "Bajo", "Cliente"),
                "CU-13": ("Alta", "Medio", "Cliente"), "CU-16": ("Crítica", "Alto", "Cliente"),
                "CU-20": ("Alta", "Medio", "Encargado"), "CU-21": ("Alta", "Medio", "Administrador"),
                "CU-22": ("Crítica", "Alto", "Cliente"), "CU-24": ("Crítica", "Alto", "Cliente"),
                "CU-25": ("Crítica", "Alto", "Cajero"), "CU-29": ("Crítica", "Alto", "Cliente")}


def estado_impl(plataforma: str) -> str:
    if "móvil" in plataforma and "Web y" in plataforma:
        return "Implementado (backend, web y móvil)"
    if "móvil" in plataforma:
        return "Implementado (backend y móvil)"
    return "Implementado (backend y web)"


def rehacer_tabla(tabla, filas_datos: list[list[str]]) -> None:
    """Deja la cabecera y reconstruye el cuerpo clonando la primera fila de datos."""
    fs = filas(tabla)
    modelo = copy.deepcopy(fs[1])
    for f in fs[1:]:
        tabla.remove(f)
    for datos in filas_datos:
        nueva = copy.deepcopy(modelo)
        for c, valor in zip(celdas(nueva), datos):
            poner_celda(c, [valor])
        tabla.append(nueva)


def buscar_tabla(inicio_cabecera: str):
    for el in body.iterchildren():
        if el.tag == qn("w:tbl") and texto(filas(el)[0]).startswith(inicio_cabecera):
            return el
    raise KeyError(inicio_cabecera)


def buscar_parrafo(inicio: str, exacto: bool = False):
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            t = texto(el).strip()
            if (t == inicio) if exacto else t.startswith(inicio):
                return el
    raise KeyError(inicio)


# =========================================================== 2. TABLAS F.T.1
# actores
t_act = buscar_tabla("IDActorDescripción")
for f in filas(t_act):
    cs = celdas(f)
    if texto(cs[0]).strip() == "A5":
        poner_celda(cs[2], ["Provee productos a la empresa. Tiene rol y cuenta en el sistema, pero en esta "
                            "versión no tiene un portal propio: sus datos y los productos que abastece los "
                            "administra el Administrador (CU-11)."])
    if texto(cs[0]).strip() == "A7":
        poner_celda(cs[2], ["Actor externo (Groq) que interpreta la búsqueda por voz y genera recomendaciones "
                            "y reportes en lenguaje natural; sus casos de uso se documentan en la "
                            "Presentación final."])

# casos de uso (Ciclos 1 y 2), ordenados por número
p_titulo_cu = buscar_parrafo("Casos de uso (Ciclo 1)", exacto=True)
poner_texto(p_titulo_cu, "Casos de uso (Ciclos 1 y 2)")
t_cu = buscar_tabla("IDCaso de UsoDescripción breve")
todos = [[i, n, d, a, pl] for (i, n, d, a, pl, _) in CICLO1]
todos += [[c["id"], c["nombre"], c["descripcion"], c["actores"], c["plataforma"]] for c in cus]
todos.sort(key=lambda x: x[0])
rehacer_tabla(t_cu, todos)

# priorización: Ciclo 1 (estado actualizado), Ciclo 2 y Ciclo 3 planificado
t_pri = buscar_tabla("IDCaso de UsoEstadoPrioridad")
rehacer_tabla(t_pri, [[i, n, estado_impl(pl), *PRIORIDAD_C1[i][:2], PRIORIDAD_C1[i][2]] for (i, n, _, _, pl, _) in CICLO1])
p_ciclo1 = buscar_parrafo("Ciclo 1", exacto=True)
poner_texto(p_ciclo1, "Ciclo 1 — Presentación 1 (flujo núcleo de punta a punta)")
t_pri2 = copy.deepcopy(t_pri)
rehacer_tabla(t_pri2, [[c["id"], c["nombre"], estado_impl(c["plataforma"]), c["prioridad"], c["riesgo"], c["actorEA"].replace(" de sucursal", "")] for c in cus])
t_pri3 = copy.deepcopy(t_pri)
rehacer_tabla(t_pri3, [
    ["CU-31", "Confirmar o rechazar transacción", "Implementado (backend: webhook firmado)", "Crítica", "Alto", "Sistema de pagos"],
    ["CU-32", "Recomendar productos al cliente", "Implementado (backend y móvil)", "Alta", "Medio", "Servicio de IA"],
    ["CU-33", "Consultar reportes de ventas e inventario", "Implementado (backend y web)", "Alta", "Bajo", "Administrador"],
    ["CU-34", "Visualizar indicadores empresariales", "Implementado (backend y web)", "Alta", "Bajo", "Administrador"],
    ["CU-35", "Recuperar contraseña", "Implementado (backend y web)", "Media", "Medio", "Todos"],
    ["CU-36", "Gestionar favoritos", "Implementado (backend y móvil)", "Baja", "Bajo", "Cliente"],
    ["CU-37", "Buscar prendas por comando de voz", "Implementado (backend y móvil)", "Media", "Medio", "Cliente"],
    ["CU-38", "Generar reporte por comando de voz", "Implementado (backend)", "Media", "Medio", "Administrador"],
    ["CU-39", "Generar prueba realista con IA", "Implementado (backend y móvil)", "Baja", "Alto", "Cliente"],
    ["CU-40", "Registrar devolución", "Implementado (backend y web)", "Media", "Medio", "Cajero"],
])
vacio = copy.deepcopy(p_ciclo1.getprevious()) if p_ciclo1.getprevious() is not None else None
p_c2 = parrafo_como(p_ciclo1, "Ciclo 2 — Presentación 2 (esta entrega)")
p_c2_nota = parrafo_como(p_ciclo1, "Criterio: se priorizaron los casos de uso que completan la operación diaria "
                                   "alrededor del núcleo del Ciclo 1 —administración del back office, entrada de "
                                   "stock por proveedores, el ciclo de la reserva visto por el cliente y por la "
                                   "sucursal, el carrito y el cobro en caja— y que ya tienen backend y pantalla, "
                                   "para que cada uno pueda demostrarse en la presentación.", negrita=False)
p_c3 = parrafo_como(p_ciclo1, "Ciclo 3 — Presentación final (planificado)")
p_c3_nota = parrafo_como(p_ciclo1, "Los CU-31 a CU-34 completan el catálogo de 34 casos de uso del enunciado. Los "
                                   "CU-35 a CU-40 son funcionalidades ya implementadas que no figuraban en el "
                                   "catálogo y que se incorporan con trazabilidad al enunciado (el reporte por voz "
                                   "y las devoluciones están pedidos explícitamente), para cerrar el proyecto en 40 "
                                   "casos de uso. Entregas y envíos a domicilio está implementado en el código pero "
                                   "queda fuera del alcance documentado, porque el enunciado solo pide los deliverys "
                                   "como tema de investigación.", negrita=False)
separador = copy.deepcopy(t_pri.getnext())  # párrafo vacío que sigue a la tabla
insertar_despues(t_pri, [copy.deepcopy(separador), p_c2, p_c2_nota, t_pri2, copy.deepcopy(separador),
                         p_c3, p_c3_nota, t_pri3])

# =========================================================== 3. FICHAS CICLO 1 (correcciones)
def tabla_ficha(cu_id: str):
    for el in body.iterchildren():
        if el.tag == qn("w:tbl") and texto(filas(el)[0]).startswith(f"CASO DE USO{cu_id} "):
            return el
    raise KeyError(cu_id)


PROTO = "Pantalla real ({}). Captura en 4.4 Prototipar interfaz de usuario."
CORRECCIONES = {
    "CU-01": {
        "FLUJO PRINCIPAL": [
            "El visitante abre el formulario de registro (web o app móvil).",
            "Ingresa nombre, apellido, correo, contraseña y, opcionalmente, teléfono y CI/NIT.",
            "El sistema valida que el correo no esté registrado.",
            "El sistema protege la contraseña con hash bcrypt.",
            "El sistema crea el usuario, le asigna el rol «cliente» y crea su registro de cliente.",
            "En la web, la aplicación inicia sesión automáticamente con las credenciales recién creadas; en la app móvil, el cliente pasa a la pantalla de inicio de sesión (CU-02)."],
        "POSTCONDICIÓN": ["Existen un usuario activo con rol «cliente» y su registro de cliente. El registro no emite token: la sesión se obtiene por CU-02."],
        "EXCEPCIÓN": ["El correo ya está registrado: el sistema rechaza el registro y sugiere iniciar sesión o recuperar la contraseña."],
        "PROTOTIPO": [PROTO.format("web /registro y app móvil")],
    },
    "CU-02": {
        "ACTORES": ["Cliente, Administrador, Encargado de sucursal, Cajero, Proveedor"],
        "FLUJO PRINCIPAL": [
            "El actor ingresa su correo electrónico y contraseña.",
            "El sistema verifica la contraseña con bcrypt; la verificación se ejecuta siempre, exista o no el correo, para no revelar por tiempo de respuesta qué correos están registrados.",
            "El sistema comprueba que el usuario esté activo.",
            "El sistema determina los roles y los permisos efectivos del usuario.",
            "El sistema emite un token de acceso y uno de refresco y registra el último acceso; la aplicación muestra la tienda al cliente y el back office al personal."],
        "EXCEPCIÓN": ["Credenciales inválidas, correo inexistente o usuario inactivo: el sistema responde siempre con el mismo mensaje genérico «Credenciales inválidas», para no permitir deducir qué cuentas existen."],
        "PROTOTIPO": [PROTO.format("web /login y app móvil")],
    },
    "CU-05": {
        "PRECONDICIÓN": ["Sesión de administrador y al menos una ciudad registrada (CU-04)."],
        "FLUJO PRINCIPAL": [
            "El administrador consulta el listado de sucursales.",
            "Crea o edita una sucursal (código, nombre, dirección, ciudad, teléfono, coordenadas y si es depósito).",
            "Define los horarios de atención por día de la semana, que luego validan las reservas (CU-16).",
            "El sistema valida que el código no se repita y que no haya dos horarios para el mismo día.",
            "El sistema guarda el cambio; si corresponde, el administrador desactiva la sucursal (borrado lógico)."],
        "EXCEPCIÓN": ["Faltan datos obligatorios o el código ya existe: el sistema rechaza el guardado.",
                      "Ya existe un horario para ese día, o la hora de cierre no es posterior a la de apertura: el sistema rechaza el horario."],
        "PROTOTIPO": [PROTO.format("web /sucursales")],
    },
    "CU-07": {
        "CASO DE USO": ["CU-07 - Gestionar catálogo maestro (categorías, tallas, colores, materiales, temporadas y colecciones)"],
        "PROPÓSITO": ["Permite al administrador mantener los catálogos maestros que estructuran el catálogo de productos: categorías, tallas, colores, materiales, temporadas y colecciones."],
        "FLUJO PRINCIPAL": [
            "Selecciona la entidad a gestionar (categoría, talla, color, material, temporada o colección).",
            "Consulta el listado de esa entidad.",
            "Crea, edita o da de baja un registro.",
            "El sistema valida que el nombre o código no se repita; en temporadas valida nombre y año, y que la fecha de fin sea posterior a la de inicio.",
            "El sistema guarda el cambio. Categorías, temporadas y colecciones se dan de baja con borrado lógico; tallas, colores y materiales se eliminan solo si ninguna variante los usa."],
        "EXCEPCIÓN": ["Se intenta desactivar una categoría con productos o subcategorías activas: el sistema rechaza la baja.",
                      "Se intenta eliminar una talla, un color o un material en uso: el sistema responde que el registro está en uso y no lo elimina.",
                      "Ya existe una temporada con el mismo nombre y año: el sistema rechaza el guardado."],
        "PROTOTIPO": [PROTO.format("web /categorias, /tallas, /colores, /temporadas y /colecciones")],
    },
    "CU-08": {
        "PRECONDICIÓN": ["Existen la categoría, la temporada, las tallas y los colores a asociar (CU-07)."],
        "FLUJO PRINCIPAL": [
            "Crea o edita un producto (nombre, descripción, categoría, material, temporada, colección, género y precio base).",
            "Elige tallas y colores; el sistema genera las variantes talla × color con su SKU y omite las combinaciones que ya existían.",
            "Opcionalmente ajusta el precio de cada variante y carga la tabla de medidas por talla.",
            "Sube una o más imágenes (JPEG, PNG o WEBP, hasta 10 MB) y marca una como principal."],
        "POSTCONDICIÓN": ["El producto queda activo con sus variantes, su tabla de medidas y sus imágenes. El registro de stock de cada variante se crea en la sucursal con el primer movimiento de inventario (recepción, CU-12)."],
        "EXCEPCIÓN": ["La combinación talla-color ya existe para el producto: el sistema no la duplica.",
                      "El archivo no es una imagen JPEG, PNG o WEBP, o supera los 10 MB: el sistema lo rechaza."],
        "PROTOTIPO": [PROTO.format("web /productos")],
    },
    "CU-09": {
        "FLUJO PRINCIPAL": [
            "El cliente abre el catálogo.",
            "El sistema lista los productos activos, paginados, con su imagen principal y su precio vigente.",
            "El cliente navega por categoría o temporada y abre el detalle de una prenda; la búsqueda con filtros se detalla en CU-10."],
        "EXCEPCIÓN": ["Catálogo vacío para el filtro elegido: el sistema muestra un estado vacío con sugerencia de quitar filtros.",
                      "Se supera el límite de consultas por minuto del catálogo público: el sistema pide reintentar."],
        "PROTOTIPO": [PROTO.format("web /catalogo y app móvil")],
    },
    "CU-13": {
        "FLUJO PRINCIPAL": [
            "El cliente elige talla y color en el detalle de la prenda.",
            "El sistema devuelve, por cada sucursal con stock registrado de esa variante, la cantidad disponible (física menos reservada).",
            "El cliente elige la sucursal donde reservar o comprar."],
        "EXCEPCIÓN": ["Ninguna sucursal tiene stock disponible de la variante: la lista llega vacía o en cero y la pantalla no ofrece sucursal para reservar."],
        "PROTOTIPO": [PROTO.format("detalle de prenda, web /producto/:id y app móvil")],
    },
    "CU-16": {
        "FLUJO PRINCIPAL": [
            "El cliente selecciona una o varias variantes.",
            "Elige la sucursal, la fecha y una franja horaria.",
            "El sistema valida que la sucursal atienda ese día y que la franja esté dentro de su horario de atención.",
            "El sistema reserva el stock de cada variante con bloqueo de fila y crea la reserva «pendiente» con su código y su fecha de expiración.",
            "El sistema notifica la reserva a los empleados de la sucursal y la confirma al cliente."],
        "POSTCONDICIÓN": ["La reserva existe en estado «pendiente», con la cantidad reservada bloqueada en el stock de cada variante."],
        "EXCEPCIÓN": ["Alguna variante no tiene stock disponible suficiente al confirmar: el sistema rechaza la reserva completa, no bloquea ninguna unidad y el cliente ajusta su selección.",
                      "La sucursal no atiende ese día o la franja está fuera de su horario: el sistema rechaza la reserva."],
        "PROTOTIPO": [PROTO.format("web /reserva/confirmar y app móvil")],
    },
    "CU-20": {
        "FLUJO PRINCIPAL": [
            "El encargado separa físicamente las prendas y cambia la reserva a «preparada».",
            "Cuando el cliente llega, confirma la llegada y la reserva pasa a «en_prueba».",
            "El cliente se prueba las prendas y decide cuáles compra.",
            "El encargado registra la decisión de cada prenda; las no seleccionadas se liberan al stock disponible.",
            "Cuando todas las prendas tienen decisión, la reserva queda «completada» y las seleccionadas pueden facturarse en caja (CU-25)."],
        "EXCEPCIÓN": ["El cliente no compra ninguna prenda: todas las unidades se liberan al stock disponible.",
                      "El cliente no se presenta: la reserva «pendiente» o «preparada» pasa a «expirada» 24 horas después del fin de la franja, por la tarea programada de expiración, y libera el stock."],
        "PROTOTIPO": [PROTO.format("web /reservas")],
    },
    "CU-21": {
        "PRECONDICIÓN": ["Existe la variante y su categoría está dentro del alcance del probador (prendas superiores masculinas)."],
        "FLUJO PRINCIPAL": [
            "El administrador sube la imagen de la prenda en PNG con transparencia real (mínimo 512 px, hasta 3 MB).",
            "Marca los puntos de referencia corporal (hombros, cadera) sobre la imagen.",
            "Ajusta el escalado y el posicionamiento de la prenda.",
            "El sistema guarda las coordenadas normalizadas (entre 0 y 1) y el administrador valida el asset."],
        "EXCEPCIÓN": ["La imagen no es un PNG con canal alfa, mide menos de 512 px o supera 3 MB: el sistema rechaza la carga.",
                      "Se intenta validar un asset sin anclajes: el sistema no lo habilita."],
        "PROTOTIPO": [PROTO.format("web /probador, editor de anclajes")],
    },
    "CU-22": {
        "FLUJO PRINCIPAL": [
            "El cliente abre el vestidor virtual en la app móvil.",
            "La app activa la cámara frontal y detecta los puntos de referencia corporal del cliente.",
            "El sistema superpone la prenda en tiempo real sobre la imagen de la cámara, ajustada a la posición y proporción detectadas.",
            "El cliente cambia de prenda desde el selector del probador, que solo ofrece las prendas con un asset validado (CU-21)."],
        "EXCEPCIÓN": ["La app no detecta la pose del cliente: deja de dibujar la prenda y muestra el mensaje «Acércate a la cámara» hasta volver a detectarla."],
        "PROTOTIPO": [PROTO.format("app móvil, pantalla Probador en modo Espejo")],
    },
    "CU-24": {
        "FLUJO PRINCIPAL": [
            "El cliente inicia el checkout desde el carrito (CU-23), en la web o en la app móvil.",
            "Elige retirar la compra en sucursal y el método de pago por pasarela (el envío a domicilio está implementado, pero queda fuera del alcance documentado).",
            "El sistema crea la venta con canal «digital» en estado «pendiente de pago», reserva el stock de cada línea y vacía el carrito.",
            "El cliente completa el pago mediante la pasarela (CU-29).",
            "Al aprobarse el pago, el sistema marca la venta «pagada» y descuenta del stock físico lo reservado; el comprobante queda disponible en el historial (CU-26)."],
        "POSTCONDICIÓN": ["La venta queda pagada, el inventario actualizado y el comprobante disponible."],
        "EXCEPCIÓN": ["El pago es rechazado: la venta pasa a «anulada», el stock reservado se libera y el cliente vuelve a armar el carrito para reintentar.",
                      "Alguna línea no tiene stock disponible al crear la venta: el sistema rechaza el checkout."],
        "PROTOTIPO": [PROTO.format("web /checkout y app móvil")],
    },
    "CU-25": {
        "PRECONDICIÓN": ["Sesión de cajero registrado como empleado de una sucursal; existe disponibilidad física de las variantes vendidas."],
        "FLUJO PRINCIPAL": [
            "El cajero agrega las variantes vendidas (búsqueda por código de barras o SKU) o carga una reserva completada.",
            "El sistema calcula subtotal, descuentos por promociones vigentes y total.",
            "El sistema registra la venta con canal «presencial» en estado «pendiente de pago».",
            "El cajero cobra en el punto de caja (CU-30); al aprobarse el cobro, la venta queda «pagada» y el stock se descuenta.",
            "El comprobante queda disponible."],
        "EXCEPCIÓN": ["No hay stock disponible suficiente de alguna variante: el sistema rechaza la venta.",
                      "La reserva cargada no está «completada» o no tiene prendas seleccionadas: el sistema no la factura."],
        "PROTOTIPO": [PROTO.format("web /caja")],
    },
    "CU-29": {
        "FLUJO PRINCIPAL": [
            "El cliente elige el método de pago por pasarela (Libélula o PayPal).",
            "El sistema verifica que la venta esté «pendiente de pago» y que no haya otro pago en curso, y crea el pago «iniciado».",
            "El sistema abre el flujo de la pasarela con el monto de la venta.",
            "El cliente completa el pago en la pasarela.",
            "La pasarela notifica el resultado por webhook firmado, o la aplicación consulta el estado; el sistema marca el pago «aprobado» o «rechazado»."],
        "POSTCONDICIÓN": ["El pago queda resuelto y asociado a la venta; si fue aprobado, la venta se confirma (CU-24). PayPal usa su sandbox real; Libélula, un simulador de sandbox."],
        "EXCEPCIÓN": ["El cliente abandona el flujo de pago: el pago queda «iniciado» y la venta «pendiente de pago» hasta un nuevo intento.",
                      "Ya hay un pago en curso para la venta: el sistema no inicia otro."],
        "PROTOTIPO": [PROTO.format("web /checkout/pago y app móvil")],
    },
}
for cu_id, campos in CORRECCIONES.items():
    poner_ficha(tabla_ficha(cu_id), campos)
p = buscar_parrafo("Diagrama de caso de uso particular — CU-07")
poner_texto(p, "Diagrama de caso de uso particular — CU-07 - Gestionar catálogo maestro (categorías, tallas, colores, materiales, temporadas y colecciones)")
for cu_id, archivo in (("CU-07", "UCP-07.png"), ("CU-08", "UCP-08.png"), ("CU-16", "UCP-16.png"), ("CU-24", "UCP-24.png")):
    reemplazar_imagen(tabla_ficha(cu_id).getprevious(), DIAG / archivo)
print("Fichas del Ciclo 1 corregidas")

# =========================================================== 4. FICHAS CICLO 2
t29 = tabla_ficha("CU-29")
titulo29 = t29.getprevious().getprevious()
vacio29 = t29.getnext()
nuevos = []
titulo_c2 = parrafo_como(titulo29, "Casos de uso del Ciclo 2")
nuevos.append(titulo_c2)
for c in cus:
    num = c["id"][3:]
    h = parrafo_como(titulo29, f"Diagrama de caso de uso particular — {c['id']} - {c['nombre']}")
    img = parrafo_imagen(DIAG / f"UCP-{num}.png")
    t = copy.deepcopy(t29)
    poner_ficha(t, {
        "CASO DE USO": [f"{c['id']} - {c['nombre']}"],
        "PROPÓSITO": [c["proposito"]],
        "ACTORES": [c["actores"]],
        "ACTOR INICIADOR": [c["iniciador"]],
        "PRECONDICIÓN": [c["pre"]],
        "FLUJO PRINCIPAL": c["flujo"],
        "POSTCONDICIÓN": [c["post"]],
        "EXCEPCIÓN": c["excepcion"],
        "PROTOTIPO": [PROTO.format(c["pantalla"])],
    })
    renumerar_flujo(t)
    nuevos += [h, img, t, copy.deepcopy(vacio29)]
insertar_despues(vacio29, nuevos)
print("16 fichas del Ciclo 2 insertadas")

# =========================================================== 5. DIAGRAMAS
# modelo de negocio: ACT-02 corregido
reemplazar_imagen(buscar_parrafo("ACT-02 — Flujo de venta y pago").getprevious(), DIAG / "ACT-02.png")

# 4.5 estructura del modelo: UC-02
cap = buscar_parrafo("UC-01 — Diagrama de casos de uso del Ciclo 1")
reemplazar_imagen(cap.getprevious(), DIAG / "UC-02.png")
poner_texto(cap, "UC-02 — Diagrama de casos de uso de los Ciclos 1 y 2 (30 casos de uso)")

# 5.1 identificar paquetes
cap = buscar_parrafo("PKG-00 – Identificar paquetes")
reemplazar_imagen(cap.getprevious(), DIAG / "PKG-00.png")
t_pk = buscar_tabla("PaqueteResponsabilidad")
fila_modelo = copy.deepcopy(filas(t_pk)[-1])
poner_celda(celdas(fila_modelo)[0], ["abastecimiento"])
poner_celda(celdas(fila_modelo)[1], ["Proveedores, productos del proveedor, órdenes de compra y recepción de mercadería (entrada de stock con costo). Se incorpora en el Ciclo 2."])
t_pk.append(fila_modelo)

# 5.1 relacionar paquetes y casos de uso: PKG-03 + filas nuevas
t_rel = buscar_tabla("Caso de usoPaquete dueño")
modelo = filas(t_rel)[-1]
for c in cus:
    f = copy.deepcopy(modelo)
    poner_celda(celdas(f)[0], [f"{c['id']} {c['nombre']}"])
    poner_celda(celdas(f)[1], [c["paquete"]])
    t_rel.append(f)
cap_pkg02 = buscar_parrafo("PKG-02 – Relacionar paquetes y casos de uso")
poner_texto(cap_pkg02, "PKG-02 – Relacionar paquetes y casos de uso (Ciclo 1): cada paquete conectado con los CU de los que es dueño")
cap3 = parrafo_como(cap_pkg02, "PKG-03 – Relacionar paquetes y casos de uso (Ciclo 2)")
insertar_despues(t_rel, [copy.deepcopy(t_rel.getnext()), parrafo_imagen(DIAG / "PKG-03.png", alto_max_in=7.5), cap3])

# vistas de paquete
for nombre, archivo in (("VP-01", "VP-01.png"), ("VP-02", "VP-02.png"), ("VP-03", "VP-03.png"), ("VP-04", "VP-04.png"),
                        ("VP-05", "VP-05.png"), ("VP-07", "VP-07.png"), ("VP-08", "VP-08.png")):
    etiqueta = buscar_parrafo(nombre + " –")
    reemplazar_imagen(etiqueta.getnext(), DIAG / archivo)
vp08 = buscar_parrafo("VP-08 –")
insertar_despues(vp08.getnext(), [parrafo_como(vp08, "VP-09 – Abastecimiento y Proveedores"), parrafo_imagen(DIAG / "VP-09.png")])


# análisis y diseño de casos de uso: COM, CLA y SEC de los 16 CU
def agregar_diagramas_cu(ultimo_titulo_prefijo: str, prefijo_archivo: str, apartado: str):
    # busca la última entrada del Ciclo 1 (CU-29) dentro del apartado
    en_apartado = False
    ref = None
    for el in body.iterchildren():
        t = texto(el).strip()
        if el.tag == qn("w:p") and t.startswith(apartado):
            en_apartado = True
        if en_apartado and el.tag == qn("w:p") and t.startswith(ultimo_titulo_prefijo):
            ref = el
            break
    titulo = ref
    imagen = ref.getnext()
    elementos = []
    for c in cus:
        num = c["id"][3:]
        elementos.append(parrafo_como(titulo, f"{c['id']} — {c['nombre']}"))
        elementos.append(parrafo_imagen(DIAG / f"{prefijo_archivo}-{num}.png"))
    insertar_despues(imagen, elementos)


agregar_diagramas_cu("CU-29 — Pagar mediante pasarela digital", "COM", "5.2 Análisis de caso de uso")
agregar_diagramas_cu("CU-29 — Pagar mediante pasarela digital", "CLA", "5.3 Análisis de clase")
agregar_diagramas_cu("CU-29 — Pagar mediante pasarela digital", "SEC", "6.2) Diseño de casos de uso")

# 5.4 análisis de paquete
cap = buscar_parrafo("PKG-01 — Paquetes involucrados en el Ciclo 1")
reemplazar_imagen(cap.getprevious(), DIAG / "PKG-01.png")
poner_texto(cap, "PKG-01 — Paquetes involucrados en los Ciclos 1 y 2 y sus dependencias")
t_dep = buscar_tabla("PaqueteDepende deCohesión")
for f in filas(t_dep):
    cs = celdas(f)
    if texto(cs[0]).strip() == "core":
        poner_celda(cs[1], ["— (ver nota)"])
        poner_celda(cs[2], ["Infraestructura transversal; única dependencia común, por diseño. Nota: core/security.py "
                            "resuelve los permisos del usuario llamando a seguridad.service con importaciones "
                            "diferidas; no consulta las tablas de seguridad, pero sí existe esa dependencia de "
                            "ejecución, que se mantiene para no duplicar la lógica de permisos."])
nueva = copy.deepcopy(filas(t_dep)[-1])
poner_celda(celdas(nueva)[0], ["abastecimiento"])
poner_celda(celdas(nueva)[1], ["catalogo, inventario, organizacion"])
poner_celda(celdas(nueva)[2], ["Alta: toda entrada de mercadería pasa por aquí y llega al stock solo a través de inventario.service.registrar_movimiento()."])
t_dep.append(nueva)

# 6.1 despliegue
cap = buscar_parrafo("DESP-01 — Arquitectura física de FashionStore")
reemplazar_imagen(cap.getprevious(), DIAG / "DESP-01.png")

# 4.4 prototipos: capturas reales de las pantallas del Ciclo 2 (web desplegada)
CAPTURAS = AQUI / "capturas"
PROTOTIPOS_C2 = [
    ("CU-03_usuarios", "CU-03 — Gestionar usuarios, roles y permisos: listado de usuarios con sus roles (/usuarios)"),
    ("CU-03_roles", "CU-03 — Gestionar usuarios, roles y permisos: roles del sistema (/roles)"),
    ("CU-04_ciudades", "CU-04 — Gestionar ciudades (/ciudades)"),
    ("CU-06_empleados", "CU-06 — Gestionar empleados y asignarlos a sucursal (/empleados)"),
    ("CU-10_busqueda_filtro", "CU-10 — Buscar y filtrar prendas: texto «Polera Nike» con filtro de categoría Poleras (/catalogo)"),
    ("CU-13_detalle_disponibilidad", "CU-13 — Consultar disponibilidad por sucursal (Ciclo 1): detalle de prenda con talla y color elegidos"),
    ("CU-11_proveedores", "CU-11 — Gestionar proveedores (/proveedores)"),
    ("CU-12_recepcion", "CU-12 — Registrar recepción de mercadería: formulario y últimas recepciones (/inventario, pestaña Recepción)"),
    ("CU-14_consolidado", "CU-14 — Consultar inventario global: existencias consolidadas por sucursal (/inventario, pestaña Consolidado)"),
    ("CU-14_alertas", "CU-14 — Consultar inventario global: alertas de reposición (/inventario, pestaña Alertas)"),
    ("CU-15_kardex", "CU-15 — Registrar movimiento de inventario: kardex de una variante (/inventario, pestaña Kardex)"),
    ("CU-15_transferencias", "CU-15 — Registrar movimiento de inventario: transferencias entre sucursales (/inventario, pestaña Transferencias)"),
    ("CU-17_mis_reservas", "CU-17 y CU-18 — Consultar estado de reserva y cancelarla desde el detalle (/mis-reservas)"),
    ("CU-19_reservas_sucursal", "CU-19 — Consultar reservas de la sucursal (/reservas)"),
    ("CU-19_dashboard_reservas_hoy", "CU-19 — Resumen «Reservas de hoy» en el panel del encargado (/dashboard)"),
    ("CU-23_carrito", "CU-23 — Gestionar carrito (/carrito)"),
    ("CU-26_mis_compras", "CU-26 — Consultar historial de compras (/mis-compras)"),
    ("CU-27_promociones", "CU-27 — Gestionar promociones (/promociones)"),
    ("CU-28_ventas_de_hoy", "CU-28 — Consultar ventas de la sucursal: ventas de hoy desde caja (/caja)"),
    ("CU-30_caja", "CU-30 — Procesar pago en caja: punto de caja con una prenda cargada (/caja)"),
]
if CAPTURAS.exists():
    titulo_45 = buscar_parrafo("4.5 Estructurar el modelo de casos de uso")
    modelo_sub = buscar_parrafo("Casos de uso del Ciclo 2", exacto=True)
    modelo_cap = buscar_parrafo("UC-02 — Diagrama de casos de uso")
    bloque = [parrafo_como(modelo_sub, "Prototipos del Ciclo 2 (capturas de la web desplegada)"),
              parrafo_como(modelo_cap, "Capturas tomadas sobre https://tienda-ropa-mauve-phi.vercel.app con los usuarios de "
                                       "cada rol. Las pantallas de reservas, compras, ventas del día y promociones aparecen sin registros "
                                       "porque la base de demostración todavía no tiene operaciones de ese tipo.")]
    for archivo, pie in PROTOTIPOS_C2:
        ruta = CAPTURAS / f"{archivo}.png"
        if ruta.exists():
            bloque += [parrafo_imagen(ruta), parrafo_como(modelo_cap, pie)]
    insertar_despues(titulo_45.getprevious(), bloque)
    print(f"Prototipos del Ciclo 2 insertados: {(len(bloque) - 2) // 2}")

# =========================================================== 6. SECCIONES NUEVAS (comparación con el grupo 32)
import io
import xml.etree.ElementTree as ET

import qrcode

URL_WEB = "https://tienda-ropa-mauve-phi.vercel.app"
URL_API = "https://backend-production-0714.up.railway.app"
URL_REPO = "https://github.com/alexosinaga00/tienda_ropa"

estilo_por_id = {s.style_id: s.name for s in doc.styles}


def p_estilo(texto_p: str, estilo_id: str | None = None, negrita=False, italica=False, tam=None, centrado=False):
    p = doc.add_paragraph(style=estilo_por_id[estilo_id]) if estilo_id else doc.add_paragraph()
    r = p.add_run(texto_p)
    r.bold = negrita
    r.italic = italica
    if tam:
        from docx.shared import Pt
        r.font.size = Pt(tam)
    if centrado:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    el = p._p
    body.remove(el)
    return el


def titulo(texto_t: str, nivel: int):
    return p_estilo(texto_t, {1: "Ttulo1", 2: "Ttulo2", 3: "Ttulo3", 4: "Ttulo4"}[nivel])


def normal(texto_n: str):
    return p_estilo(texto_n)


def vineta(texto_v: str):
    return p_estilo(texto_v, "Listaconvietas")


# sombreado de cabecera tomado de la tabla de casos de uso existente
_cab_tcpr = copy.deepcopy(celdas(filas(buscar_tabla("IDCaso de UsoDescripción breve"))[0])[0].find(qn("w:tcPr")))


def tabla(cabecera: list[str], filas_t: list[list[str]], tam=9):
    from docx.shared import Pt
    t = doc.add_table(rows=1 + len(filas_t), cols=len(cabecera))
    t.style = doc.styles[estilo_por_id["Tablaconcuadrcula"]]
    for i, fila in enumerate([cabecera] + filas_t):
        for j, valor in enumerate(fila):
            celda = t.rows[i].cells[j]
            celda.text = ""
            run = celda.paragraphs[0].add_run(str(valor))
            run.font.size = Pt(tam)
            run.bold = i == 0
            if i == 0 and _cab_tcpr is not None:
                shd = _cab_tcpr.find(qn("w:shd"))
                if shd is not None:
                    celda._tc.get_or_add_tcPr().append(copy.deepcopy(shd))
    el = t._tbl
    body.remove(el)
    return el


def pie(texto_pie: str):
    return p_estilo(texto_pie, italica=True, tam=9, centrado=True)


def figura(archivo: Path, texto_pie: str, alto_max=8.0):
    return [parrafo_imagen(archivo, alto_max_in=alto_max), pie(texto_pie)]


# --- 6.1 Perfil: objetivos específicos del enunciado + RF/RNF en el alcance -------
OBJETIVOS = [
    "Analizar y especificar los requisitos funcionales y no funcionales del sistema.",
    "Gestionar clientes, usuarios, empleados, proveedores y sucursales.",
    "Administrar el catálogo de prendas, categorías, tallas, colores, temporadas y colecciones.",
    "Consultar la disponibilidad de prendas por sucursal.",
    "Permitir al cliente reservar varias prendas para probarlas posteriormente en una tienda.",
    "Gestionar el proceso de recepción y atención de reservas.",
    "Incorporar vestidores virtuales utilizando realidad aumentada.",
    "Permitir compras mediante la plataforma web y aplicación móvil.",
    "Permitir pagos presenciales en puntos de caja.",
    "Integrar una pasarela de pago para compras digitales.",
    "Actualizar automáticamente el inventario después de reservas, ventas, devoluciones y recepción de productos.",
    "Gestionar productos provenientes de diferentes proveedores y temporadas.",
    "Incorporar funcionalidades de inteligencia artificial para recomendación o asistencia al cliente.",
    "Generar reportes y dashboards para apoyar la toma de decisiones.",
    "Aplicar UML 2.5+ para modelar los procesos, funcionalidades y arquitectura del sistema.",
    "Implementar el sistema utilizando FastAPI, Angular y Flutter/Dart.",
]
tit_obj = buscar_parrafo("1.3 Objetivos Específicos", exacto=True)
objetivos_viejos = []
sig = tit_obj.getnext()
def _estilo_de(p) -> str:
    ppr = p.find(qn("w:pPr"))
    st = ppr.find(qn("w:pStyle")) if ppr is not None else None
    return st.get(qn("w:val")) if st is not None else ""


while sig is not None and sig.tag == qn("w:p") and _estilo_de(sig) == "Listaconnmeros":
    objetivos_viejos.append(sig)
    sig = sig.getnext()
modelo_obj = objetivos_viejos[0]
for p in objetivos_viejos:
    body.remove(p)
insertar_despues(tit_obj, [parrafo_como(modelo_obj, o) for o in OBJETIVOS])

req = json.loads((AQUI / "requisitos.json").read_text(encoding="utf8"))
ultimo_alcance = buscar_parrafo("Ventas y pagos: carrito con las promociones vigentes")
insertar_despues(ultimo_alcance, [
    titulo("Requisitos funcionales", 3),
    normal("Los 25 requisitos funcionales del enunciado, con los casos de uso que los realizan y el ciclo en que se documentan:"),
    tabla(["RF", "Requisito", "Casos de uso", "Ciclo"], req["rf"]),
    titulo("Requisitos no funcionales", 3),
    normal("RNF01 a RNF09 son los del enunciado; RNF10 a RNF15 se agregan porque el sistema construido los cumple y conviene declararlos. La columna de la derecha indica cómo se cumple cada uno en el código:"),
    tabla(["RNF", "Requisito", "Cómo se cumple"], req["rnf"]),
])

# --- 6.2 Marco teórico: deliverys ---------------------------------------------------
tit_24 = buscar_parrafo("2.4) UML", exacto=True)
insertar_despues(tit_24.getprevious(), [
    titulo("2.3.5 Deliverys", 3),
    normal("Un delivery es el servicio que traslada un pedido desde el comercio hasta la ubicación del cliente. En una "
           "plataforma de comercio electrónico es la etapa que sigue al pago: se confirma el pedido, el comercio lo "
           "prepara, se asigna un repartidor, este lo recoge y lo transporta, y la entrega se confirma en destino."),
    normal("Como usuario. En Bolivia, plataformas como Yango y Yummy conectan tres partes: el cliente, que elige "
           "productos, indica su ubicación y paga desde la aplicación; el comercio, que recibe y prepara el pedido; y "
           "el repartidor, que acepta el viaje y comparte su posición en tiempo real. La aplicación muestra al cliente "
           "el tiempo estimado y el estado del pedido (confirmado, en preparación, en camino, entregado)."),
    normal("Cómo se calcula el costo. No existe una fórmula única; las plataformas combinan estos factores:"),
    vineta("Distancia entre el punto de recogida y el destino, o la zona tarifaria del destino."),
    vineta("Peso y volumen del paquete, que condicionan el tipo de vehículo."),
    vineta("Demanda y frecuencia: tarifas dinámicas en horas pico y descuentos para clientes o comercios frecuentes."),
    vineta("Urgencia y franja horaria de la entrega."),
    normal("Una forma general es: costo = tarifa base de la zona + costo por distancia + recargo por peso o volumen + "
           "recargo por urgencia o demanda."),
    normal("Aplicación en FashionStore. El enunciado centra la compra en la reserva con prueba en sucursal y en la venta "
           "digital, por eso el delivery queda fuera del alcance documentado. Aun así, el backend implementa un "
           "tarifario simplificado en el paquete entregas: el costo de envío es la tarifa base de la zona (definida por "
           "anillos de la ciudad) más un recargo según el peso estimado del pedido (cantidad de prendas × peso "
           "promedio), y el envío recorre los estados programado → en ruta → entregado o fallido."),
])

# --- 6.3 Numeración de capítulos ------------------------------------------------------
for viejo, nuevo in (("F.T.1 — Captura de Requisitos", "4) F.T.1 — Captura de Requisitos"),
                     ("F.T.2 — Análisis", "5) F.T.2 — Análisis"),
                     ("F.T.3 — Diseño", "6) F.T.3 — Diseño"),
                     ("F.T.4 — Implementación", "7) F.T.4 — Implementación")):
    poner_texto(buscar_parrafo(viejo, exacto=True), nuevo)

# --- 6.4 URLs unificadas --------------------------------------------------------------
for prefijo, nuevo_txt in (("API: https://tiendaropa-production", f"API: {URL_API} — GET /health responde 200."),
                           ("Web: https://tienda-ropa-ruby.vercel.app", f"Web: {URL_WEB} — la tienda pública (/catalogo) y el back office responden 200.")):
    p = buscar_parrafo(prefijo)
    poner_texto(p, nuevo_txt)
p = buscar_parrafo("Web: " + URL_WEB)
if p.getnext() is not None and texto(p.getnext()).strip() == "responden 200.":
    body.remove(p.getnext())

# --- 6.5 Prototipos del Ciclo 1: pies de figura ----------------------------------------
PIES_C1 = ["Tienda web: catálogo público (CU-09)", "Registro de cliente (CU-01) e inicio de sesión (CU-02)",
           "Panel del back office del administrador", "Usuarios (CU-03)", "Roles (CU-03)", "Inventario consolidado (CU-14)"]
tit_proto = buscar_parrafo("Prototipar interfaz de usuario", exacto=True)
poner_texto(tit_proto, "4.4 Prototipar interfaz de usuario")
_ppr_proto = tit_proto.find(qn("w:pPr"))
if _ppr_proto is not None and _ppr_proto.find(qn("w:numPr")) is not None:
    _ppr_proto.remove(_ppr_proto.find(qn("w:numPr")))
imagenes_c1 = []
el = tit_proto.getnext()
while el is not None and texto(el).strip() != "Prototipos del Ciclo 2 (capturas de la web desplegada)":
    if el.tag == qn("w:p") and list(el.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")):
        imagenes_c1.append(el)
    el = el.getnext()
insertar_despues(tit_proto, [titulo("Prototipos del Ciclo 1", 3)])
for img_p, texto_pie in zip(imagenes_c1, PIES_C1):
    ppr = img_p.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        ppr.remove(ppr.find(qn("w:numPr")))
    img_p.addnext(pie(texto_pie))

# --- 6.6 Diseño de casos de uso: estado, tiempo y navegación --------------------------
tit_datos = buscar_parrafo("6.3) Diseño de datos")
insertar_despues(tit_datos.getprevious(), [
    titulo("Diagramas de estado", 3),
    normal("Se modelan las tres entidades cuyo ciclo de vida está controlado por el código: la reserva, la venta y el "
           "pago. Cada transición indica el evento (operación del service), la condición entre corchetes y el efecto "
           "después de la barra; las transiciones no listadas son rechazadas por el sistema."),
    *figura(DIAG / "EST-01.png", "EST-01 — Estados de la reserva (reservas/service.py: TRANSICIONES_VALIDAS y expirar_reservas)"),
    *figura(DIAG / "EST-02.png", "EST-02 — Estados de la venta. El catálogo de estados incluye «entregada», pero ningún flujo actual la asigna: el envío lleva sus propios estados."),
    *figura(DIAG / "EST-03.png", "EST-03 — Estados del pago (pasarela y caja)"),
    titulo("Diagramas de tiempo", 3),
    normal("Muestran cómo cambian de estado las entidades a lo largo del tiempo en los tres escenarios donde el tiempo "
           "es parte de la regla de negocio: la expiración de una reserva, la confirmación de un pago por pasarela y la "
           "atención de una reserva dentro de su franja. Se dibujan con la notación UML de State Lifeline (estados en el "
           "eje vertical, eventos en cada transición y restricciones de duración entre llaves)."),
    *figura(DIAG / "TIE-01.png", "TIE-01 — Reserva no retirada que expira 24 horas después del fin de su franja"),
    *figura(DIAG / "TIE-02.png", "TIE-02 — Compra digital: la venta se confirma y el stock se descuenta recién cuando la pasarela aprueba el pago"),
    *figura(DIAG / "TIE-03.png", "TIE-03 — Atención de una reserva en sucursal con compra parcial"),
    titulo("Diagramas de navegación", 3),
    normal("Recorrido de pantallas de cada aplicación, organizado por sistema: la tienda web del cliente, el back office "
           "web del personal y la aplicación móvil. Cada flecha «link» es una acción del usuario que lleva de una "
           "pantalla a otra."),
    *figura(DIAG / "NAV-01.png", "NAV-01 — Navegación de la tienda web (cliente)"),
    *figura(DIAG / "NAV-02.png", "NAV-02 — Navegación del back office web (administrador, encargado y cajero)"),
    *figura(DIAG / "NAV-03.png", "NAV-03 — Navegación de la aplicación móvil (cliente)"),
])

# --- 6.7 Diseño de datos: diagramas por dominio, mapeo, tablas de volumen -------------
cap_dcd = buscar_parrafo("Diagrama de clases de diseño (DCD)")
img_dcd = cap_dcd.getprevious()
nuevos_dat = []
for archivo, pie_txt in (("DAT-01", "DAT-01 — Seguridad y organización"), ("DAT-02", "DAT-02 — Catálogo"),
                          ("DAT-03", "DAT-03 — Inventario"), ("DAT-04", "DAT-04 — Reservas"),
                          ("DAT-05", "DAT-05 — Probador"), ("DAT-06", "DAT-06 — Ventas y pagos")):
    if (DIAG / f"{archivo}.png").exists():
        nuevos_dat += figura(DIAG / f"{archivo}.png", pie_txt)
if nuevos_dat:
    # el DCD general se conserva; los diagramas por dominio van a continuación
    insertar_despues(cap_dcd, [normal("El detalle de cada dominio se muestra en los diagramas de datos siguientes:")] + nuevos_dat)

datos = json.loads((AQUI / "datos_fisicos.json").read_text(encoding="utf8"))
filas_mapeo = []
for t in datos:
    pk = ", ".join(c["nombre"] for c in t["columnas"] if "PK" in c["llave"])
    fk = "; ".join(f"{c['nombre']} → {c['llave'].split('FK → ')[1].split(',')[0]}" for c in t["columnas"] if "FK" in c["llave"])
    filas_mapeo.append([t["nombre"], t["paquete"], pk, fk or "—", str(len(t["columnas"]))])
tit_fisico = buscar_parrafo("6.3.2) Diseño de datos físicos")
insertar_despues(tit_fisico.getprevious(), [
    titulo("Mapeo de clases a tablas", 4),
    normal("Cada clase del modelo de diseño se mapea a una tabla con el mismo nombre en singular y snake_case. La "
           "tabla siguiente resume, para las 60 tablas implementadas, su paquete dueño, su clave primaria y sus claves "
           "foráneas; las relaciones N:M (usuario_rol, rol_permiso, producto_proveedor, producto_coleccion) usan clave "
           "primaria compuesta."),
    tabla(["Tabla", "Paquete", "PK", "FK (→ tabla referenciada)", "Columnas"], filas_mapeo, tam=8),
])

bloques_volumen = [titulo("Tablas de volumen", 4),
                   normal("Detalle físico de cada tabla, tomado de los modelos SQLAlchemy del backend (los mismos "
                          "que crean las migraciones de Alembic), por lo que coincide con la base desplegada.")]
for t in datos:
    bloques_volumen.append(p_estilo(f"{t['nombre'].upper()}  ·  paquete {t['paquete']}", negrita=True, tam=9))
    bloques_volumen.append(tabla(["Atributo", "Tipo", "Tamaño", "Nulo", "Llave", "Descripción"],
                                 [[c["nombre"], c["tipo"], c["tamano"], c["nulo"], c["llave"], c["descripcion"]] for c in t["columnas"]], tam=8))
extracto = buscar_parrafo("Extracto real del esquema (producto_variante y stock)")
insertar_despues(extracto.getprevious(), bloques_volumen + [
    normal("El script completo de creación de las 60 tablas en PostgreSQL está en el Anexo C."),
])

# --- 6.8 F.T.5 Pruebas ---------------------------------------------------------------------
junit = ET.parse(AQUI / "pytest_resultado.xml").getroot()
suite = junit if junit.tag == "testsuite" else junit.find("testsuite")
casos = list(suite.iter("testcase"))
nombres_ok = {c.get("name") for c in casos if not list(c)}
por_archivo: dict[str, list[int]] = {}
for c in casos:
    modulo = c.get("classname").replace("tests.", "")
    total, ok = por_archivo.get(modulo, [0, 0])
    por_archivo[modulo] = [total + 1, ok + (0 if list(c) else 1)]

pruebas = json.loads((AQUI / "pruebas.json").read_text(encoding="utf8"))
bloques_pruebas = [
    titulo("8) F.T.5 — Pruebas", 1),
    titulo("8.1 Pruebas automatizadas del backend", 2),
    normal(f"La suite de pruebas del backend (pytest, base SQLite en memoria con los mismos modelos y seeds) se "
           f"ejecutó el 13/09/2026: {len(nombres_ok)} de {len(casos)} pruebas pasaron, en {float(suite.get('time')):.0f} "
           f"segundos. Cada prueba levanta la API completa con TestClient y verifica el comportamiento a través de HTTP, "
           f"por lo que valida router, service, repository y modelo juntos."),
    tabla(["Módulo de pruebas", "Pruebas", "Aprobadas"], [[m, str(v[0]), str(v[1])] for m, v in sorted(por_archivo.items())]),
    titulo("8.2 Pruebas de caja negra por caso de uso", 2),
    normal("Cada caso de uso de los Ciclos 1 y 2 se prueba desde afuera, por su flujo principal y sus excepciones. La "
           "columna Evidencia indica cómo se verificó cada paso: A = prueba automatizada del backend que pasó en la "
           "ejecución anterior (se cita su nombre); M = verificación manual en la web desplegada el 13/09/2026, durante "
           "la toma de las capturas del prototipo."),
]
faltantes = []
for caso in pruebas:
    filas_p = []
    for i, (accion, esperado, evidencia) in enumerate(caso["pasos"], start=1):
        if evidencia.startswith("A:"):
            nombres = [n.strip() for n in evidencia[2:].split(";")]
            faltantes += [n for n in nombres if n not in nombres_ok]
            estado = "Satisfactorio" if all(n in nombres_ok for n in nombres) else "Sin evidencia"
        elif "revisión de código" in evidencia:
            estado = "Verificado en código"
        else:
            estado = "Satisfactorio"
        filas_p.append([str(i), accion, esperado, evidencia, estado])
    resultado = "Satisfactorio" if all(f[4] == "Satisfactorio" for f in filas_p) else "Satisfactorio con observación"
    bloques_pruebas += [
        titulo(f"{caso['id']} — {caso['nombre']}", 3),
        normal("Precondiciones: " + " ".join(caso["pre"])),
        tabla(["Paso", "Acción", "Resultado esperado", "Evidencia", "Estado"], filas_p, tam=8),
        normal(f"Responsable: {caso['responsable']}. Resultado de la prueba: {resultado}."),
    ]
if faltantes:
    print("  [OJO] pruebas citadas que no existen o no pasaron:", sorted(set(faltantes)))

bloques_pruebas += [
    titulo("Conclusiones", 1),
    normal("Los dos primeros ciclos permitieron construir y documentar 30 de los 34 casos de uso del catálogo, siguiendo "
           "el PUDS de principio a fin: cada caso de uso tiene su ficha, su diagrama particular, su análisis de "
           "comunicación y de clases y su diseño de secuencia, y todos están implementados en un backend desplegado en la "
           "nube con clientes web y móvil."),
    normal("El Ciclo 1 resolvió el núcleo con mayor riesgo técnico —la concurrencia de stock entre reserva y venta, la "
           "integración de pasarelas y el vestidor virtual con detección de pose en el dispositivo— y el Ciclo 2 completó "
           "la operación diaria alrededor de ese núcleo: administración del back office, entrada de mercadería con costo "
           "promedio ponderado, ciclo completo de la reserva, carrito, promociones y cobro en caja."),
    normal("Revisar el documento contra el código en cada ciclo resultó tan importante como agregar casos de uso: varias "
           "fichas del Ciclo 1 prometían comportamientos que el sistema no tenía (una reserva parcial, un carrito que se "
           "conserva tras un pago rechazado) y se corrigieron para que el modelo describa el sistema construido. Las "
           "234 pruebas automatizadas del backend, todas aprobadas, respaldan las pruebas de caja negra de los casos de uso."),
    titulo("Recomendaciones", 1),
    vineta("Documentar en el Ciclo 3 los casos de uso restantes (confirmación de transacción, recomendador, reportes y "
           "dashboard) y las funcionalidades complementarias ya implementadas, hasta completar 40."),
    vineta("Agregar un planificador (por ejemplo, un cron de Railway) que invoque /tareas/expirar-reservas; hoy la "
           "expiración depende de una llamada externa."),
    vineta("Cerrar la deuda técnica abierta: verificar en pagar_en_caja que no exista un pago de pasarela en curso, "
           "aplicar el permiso probador.usar y devolver la disponibilidad en el detalle de producto."),
    vineta("Cargar datos de demostración de reservas, ventas y promociones antes de la defensa, para que las pantallas "
           "de consulta muestren información real."),
    vineta("Agregar pruebas automatizadas para los casos de uso que hoy solo tienen verificación manual (historial de "
           "compras y ventas de la sucursal)."),
]
bibliografia = buscar_parrafo("Bibliografía", exacto=True)
insertar_despues(bibliografia.getprevious(), bloques_pruebas)

# --- 6.9 Anexos: URLs con QR y script SQL ----------------------------------------------
def qr_parrafo(url: str):
    img = qrcode.make(url)
    ruta = AQUI / "diagramas" / ("QR-" + url.split("//")[1].split("/")[0].replace(".", "_") + ".png")
    img.save(ruta)
    return parrafo_imagen(ruta, ancho_max_in=1.6, alto_max_in=1.6)


# se quitan los QR y el enlace viejos (apuntaban al despliegue anterior)
for el in list(body.iterchildren()):
    t = texto(el)
    if el.tag == qn("w:p") and ("Qr del frontend" in t or t.strip() == "https://tienda-ropa-ruby.vercel.app/catalogo"):
        body.remove(el)

glosario = buscar_tabla("SiglaSignificado")
anexos = [
    titulo("Anexo B — Repositorio y despliegue", 2),
    tabla(["Recurso", "URL"], [["Web (tienda y back office)", URL_WEB], ["API (Swagger en /docs)", URL_API],
                               ["Repositorio (backend, web y móvil)", URL_REPO]]),
    normal("Los códigos QR de estos enlaces se encuentran en el Anexo D."),
    titulo("Anexo C — Script de creación de la base de datos (PostgreSQL)", 2),
    normal("Script DDL de PostgreSQL correspondiente a los modelos del backend; es equivalente al resultado de aplicar "
           "las 17 migraciones de Alembic."),
]
for bloque_sql in (AQUI / "script_postgres.sql").read_text(encoding="utf8").split("\n\n"):
    p = doc.add_paragraph()
    from docx.shared import Pt
    for k, linea in enumerate(bloque_sql.strip("\n").split("\n")):
        run = p.add_run(linea.replace("\t", "    "))
        run.font.name = "Consolas"
        run.font.size = Pt(7.5)
        if k < len(bloque_sql.strip("\n").split("\n")) - 1:
            run.add_break()
    body.remove(p._p)
    anexos.append(p._p)
# anexo final: códigos QR
anexos += [
    titulo("Anexo D — Códigos QR del sistema", 2),
    p_estilo("Aplicación web (tienda y back office)", negrita=True, centrado=True), qr_parrafo(URL_WEB),
    pie(URL_WEB),
    p_estilo("API del backend (documentación Swagger)", negrita=True, centrado=True), qr_parrafo(URL_API + "/docs"),
    pie(URL_API + "/docs"),
    p_estilo("Repositorio del código fuente", negrita=True, centrado=True), qr_parrafo(URL_REPO),
    pie(URL_REPO),
]
insertar_despues(glosario.getnext() if glosario.getnext() is not None else glosario, anexos)

# limpieza final: marcas de formato Markdown que quedaron como texto literal
for t_el in body.iter(qn("w:t")):
    if t_el.text and "**" in t_el.text:
        t_el.text = t_el.text.replace("**", "")
print("Secciones nuevas agregadas")

# =========================================================== 7. CORRECCIONES DE LA REVISIÓN
import re


def reemplazar_en_textos(viejo: str, nuevo: str) -> None:
    n = 0
    for t_el in body.iter(qn("w:t")):
        if t_el.text and viejo in t_el.text:
            t_el.text = t_el.text.replace(viejo, nuevo)
            n += 1
    if n == 0:
        print("  [OJO] sin coincidencias:", viejo[:60])


def estilo(p) -> str:
    ppr = p.find(qn("w:pPr"))
    st = ppr.find(qn("w:pStyle")) if ppr is not None else None
    return st.get(qn("w:val")) if st is not None else ""


def cambiar_estilo(p, estilo_id: str) -> None:
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = p.makeelement(qn("w:pPr"), {})
        p.insert(0, ppr)
    st = ppr.find(qn("w:pStyle"))
    if st is None:
        st = ppr.makeelement(qn("w:pStyle"), {})
        ppr.insert(0, st)
    st.set(qn("w:val"), estilo_id)
    for r in p.findall(qn("w:r")):  # el estilo del título manda sobre la negrita del run
        rpr = r.find(qn("w:rPr"))
        if rpr is not None and rpr.find(qn("w:b")) is not None:
            rpr.remove(rpr.find(qn("w:b")))


# 2.2.4: ya no hay contenido pendiente para la final salvo el Ciclo 3
poner_texto(buscar_parrafo("Esta presentación cubre Captura de Requisitos"),
            "Esta presentación cubre, para los Ciclos 1 y 2, los flujos de Captura de Requisitos, Análisis, Diseño "
            "(incluidos los diagramas de estado, tiempo y navegación y el diseño de datos), Implementación y Pruebas. La "
            "Presentación final incorpora los 4 casos de uso restantes del catálogo y las funcionalidades complementarias "
            "ya implementadas.")

# paquetes y capas coherentes con 7.2 (12 paquetes de negocio + core, cinco archivos por paquete)
poner_texto(buscar_parrafo("En este proyecto, la arquitectura se diseña como un Monolito Modular"),
            "En este proyecto, la arquitectura se diseña como un Monolito Modular estructurado sobre FastAPI (Python). Este "
            "enfoque permite encapsular los doce paquetes de negocio (seguridad, organización, catálogo, abastecimiento, "
            "inventario, reservas, ventas, pagos, entregas, probador, inteligencia y reportes) y el núcleo core en una misma "
            "base de código, con transacciones atómicas de base de datos y despliegue simplificado en una plataforma PaaS "
            "(Railway), evitando la sobrecarga de una arquitectura de microservicios para un proyecto de cuatro semanas.")
poner_texto(buscar_parrafo("Cada paquete de negocio (seguridad, organización"),
            "Cada paquete de negocio tiene los mismos cinco archivos, y cada uno pertenece a una capa: router.py "
            "(presentación HTTP), schemas.py (contrato de entrada y salida), service.py (reglas de negocio), repository.py "
            "(acceso a datos, heredando de core/crud_base.py) y models.py (persistencia con SQLAlchemy). La comunicación entre "
            "paquetes ocurre siempre service a service: ningún paquete consulta directamente las tablas de otro.")
poner_texto(buscar_parrafo("PKGC-01"),
            "PKGC-01 — Capas de cada paquete de negocio: presentación (router), negocio (service), acceso a datos (repository) "
            "y persistencia (models); los esquemas de schemas.py son el contrato entre presentación y negocio")
reemplazar_en_textos("se registran los routers de los trece paquetes", "se registran los routers de los doce paquetes de negocio y del núcleo core")

# mapeo: producto_coleccion no existe
reemplazar_en_textos("producto_proveedor, producto_coleccion)", "producto_proveedor, favorito)")

# pie de PKG-02
poner_texto(buscar_parrafo("PKG-02 – Relacionar paquetes y casos de uso"),
            "PKG-02 – Relacionar paquetes y casos de uso: el diagrama corresponde al Ciclo 1 y la tabla incluye los casos de uso de los Ciclos 1 y 2")

# DCD: se conserva la imagen original del documento (no se reemplaza)

# deliverys: ejemplo del enunciado y bibliografía
reemplazar_en_textos("plataformas como Yango y Yummy", "plataformas como Yaigo y Yummy")
ultima_ref = buscar_parrafo("Stripe. (s.f.). Stripe Documentation.")
insertar_despues(ultima_ref, [parrafo_como(ultima_ref, r) for r in (
    "Google Cloud. (s.f.). Vertex AI Documentation. https://cloud.google.com/vertex-ai/docs",
    "Sparx Systems. (s.f.). Enterprise Architect User Guide. https://sparxsystems.com/enterprise_architect_user_guide/",
    "Vercel. (s.f.). Vercel Documentation. https://vercel.com/docs",
    "Yaigo. (s.f.). Yaigo: aplicación de delivery y transporte en Bolivia.",
    "Yummy. (s.f.). Yummy: aplicación de delivery en Bolivia.",
)])

# 2.4.3: todos los diagramas usados
vineta_paquetes = buscar_parrafo("Diagrama de paquetes: organiza el sistema en módulos")
insertar_despues(vineta_paquetes, [parrafo_como(vineta_paquetes, v) for v in (
    "Diagrama de actividades: modela los procesos del negocio (reserva y prueba en sucursal, venta y pago) antes de detallar los casos de uso.",
    "Diagrama de estados: describe los estados de una entidad (reserva, venta, pago) y los eventos que provocan cada transición.",
    "Diagrama de tiempo: muestra cómo cambia el estado de las entidades a lo largo del tiempo cuando el tiempo es parte de la regla de negocio, como la expiración de una reserva.",
    "Diagrama de navegación: representa el recorrido de pantallas de cada aplicación.",
    "Diagrama de componentes: muestra los paquetes del backend como componentes y sus dependencias reales.",
)])

# numeración de títulos: 2.1 con su 2.1.1, y subniveles sin paréntesis
t21 = buscar_parrafo("2.1 Comercio electrónico multi-sucursal: fundamentos", exacto=True)
poner_texto(t21, "2.1 Comercio electrónico multi-sucursal")
insertar_despues(t21, [titulo("2.1.1 Fundamentos", 3)])
for nombre_t in ("2.1.2 Generalidades", "2.1.3 Características", "2.1.4 Conceptos específicos del rubro y la arquitectura"):
    cambiar_estilo(buscar_parrafo(nombre_t, exacto=True), "Ttulo3")
for p in body.iterchildren():
    if p.tag == qn("w:p") and estilo(p) in ("Ttulo2", "Ttulo3", "Ttulo4"):
        t = texto(p)
        nuevo_t = re.sub(r"^(\d+(?:\.\d+)+)\)\s*", r"\1 ", t)
        if nuevo_t != t:
            poner_texto(p, nuevo_t)
poner_texto(buscar_parrafo("Diagramas De Secuencia", exacto=True), "Diagramas de secuencia")

# 4.3 y 4.4: subtítulos de ciclo como títulos, prototipos sin duplicados ni viñetas vacías
t43 = buscar_parrafo("4.3 Detallar casos de uso", exacto=True)
insertar_despues(t43, [titulo("Casos de uso del Ciclo 1", 3)])
cambiar_estilo(buscar_parrafo("Casos de uso del Ciclo 2", exacto=True), "Ttulo3")
p_c2proto = buscar_parrafo("Prototipos del Ciclo 2 (capturas de la web desplegada)", exacto=True)
cambiar_estilo(p_c2proto, "Ttulo3")
el = buscar_parrafo("Prototipos del Ciclo 1", exacto=True).getnext()
while el is not None and el is not p_c2proto:
    sig = el.getnext()
    t = texto(el).strip()
    tiene_img = bool(list(el.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")))
    if el.tag == qn("w:p") and not t and not tiene_img:
        body.remove(el)
    elif el.tag == qn("w:p") and tiene_img and estilo(el) == "Listaconnmeros":
        # la imagen heredaba la numeración del estilo de lista ("7." suelto)
        ppr = el.find(qn("w:pPr"))
        ppr.remove(ppr.find(qn("w:pStyle")))
        if ppr.find(qn("w:jc")) is None:
            ppr.append(ppr.makeelement(qn("w:jc"), {qn("w:val"): "center"}))
    elif t in ("Usuarios (CU-03)", "Roles (CU-03)", "Inventario consolidado (CU-14)"):
        body.remove(el.getprevious())
        body.remove(el)
    el = sig

# 5.2, 5.3 y 6.2: separadores de ciclo
for apartado in ("5.2 Análisis de caso de uso", "5.3 Análisis de clase", "Diagramas de secuencia"):
    dentro = False
    for p in list(body.iterchildren()):
        t = texto(p).strip()
        if p.tag == qn("w:p") and t.startswith(apartado):
            dentro = True
            continue
        if not dentro or estilo(p) != "Ttulo4":
            continue
        if t.startswith("CU-01 —"):
            p.addprevious(p_estilo("Casos de uso del Ciclo 1", negrita=True))
        elif t.startswith("CU-03 —"):
            p.addprevious(p_estilo("Casos de uso del Ciclo 2", negrita=True))
            break

# F.T.4: unir los renglones que quedaron como párrafos separados
inicio = buscar_parrafo("7) F.T.4 — Implementación", exacto=True)
fin = buscar_parrafo("8) F.T.5 — Pruebas", exacto=True)


def es_parrafo_plano(p) -> bool:
    return (p is not None and p.tag == qn("w:p") and estilo(p) == "" and texto(p).strip() != ""
            and not list(p.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")))


p = inicio.getnext()
while p is not None and p is not fin:
    sig = p.getnext()
    while es_parrafo_plano(p) and es_parrafo_plano(sig) and not texto(p).rstrip().endswith((".", ":", ";", "?", "!")):
        ts = list(p.iter(qn("w:t")))
        if ts:
            ts[-1].text = (ts[-1].text or "").rstrip() + " "
            ts[-1].set(qn("xml:space"), "preserve")
        for r in sig.findall(qn("w:r")):
            p.append(r)
        body.remove(sig)
        sig = p.getnext()
    p = sig

# portada y comillas
for t_el in body.iter(qn("w:t")):
    if t_el.text and t_el.text.strip() == "Aguayo Quiroz Luis Migue":
        t_el.text = t_el.text.replace("Luis Migue", "Luis Miguel")

# glosario ampliado
t_glos = buscar_tabla("SiglaSignificado")
modelo_glos = filas(t_glos)[-1]
for sigla, significado in (
    ("ACID", "Atomicidad, Consistencia, Aislamiento y Durabilidad"),
    ("B2C", "Business to Consumer (de empresa a consumidor)"),
    ("CORS", "Cross-Origin Resource Sharing (intercambio de recursos entre orígenes)"),
    ("DDL", "Data Definition Language (lenguaje de definición de datos)"),
    ("EA", "Enterprise Architect (herramienta de modelado UML)"),
    ("FK", "Foreign Key (llave foránea)"),
    ("HTTPS", "HyperText Transfer Protocol Secure"),
    ("JWT", "JSON Web Token"),
    ("MVP", "Minimum Viable Product (producto mínimo viable)"),
    ("ORM", "Object-Relational Mapping (mapeo objeto-relacional)"),
    ("QR", "Quick Response (código de respuesta rápida)"),
    ("RA", "Realidad aumentada"),
    ("REST", "Representational State Transfer"),
    ("RF", "Requisito funcional"),
    ("RNF", "Requisito no funcional"),
    ("SKU", "Stock Keeping Unit (código de inventario de la variante)"),
):
    fila = copy.deepcopy(modelo_glos)
    poner_celda(celdas(fila)[0], [sigla])
    poner_celda(celdas(fila)[1], [significado])
    t_glos.append(fila)
print("Correcciones de la revisión aplicadas")

# limpieza del párrafo auxiliar
body.remove(_scratch._p)
doc.save(str(DESTINO))
print("Documento generado:", DESTINO)
